"""Excel export of one recorded run (samples + events) to an .xlsx workbook."""

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import storage

COLUMNS = [
    "Run ID", "Protocol Name", "Protocol Version", "Timestamp", "Elapsed Time",
    "Process Step", "Actual Temperature (°C)", "Target Temperature (°C)",
    "Temperature Deviation (°C)", "Cooling Rate (°C/min)", "Relative Humidity (% RH)",
    "System Status", "Event Type", "Event Details",
]


def _hms(sec):
    if sec is None:
        return None
    sec = int(round(sec))
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def _ts(iso):
    try:
        return datetime.fromisoformat(iso)
    except (TypeError, ValueError):
        return iso


def build_run_workbook(run_id):
    """Return xlsx bytes, or None when the run does not exist."""
    run = storage.get_run(run_id)
    if not run:
        return None
    samples = storage.run_samples(run_id)
    events = storage.run_events(run_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Run Data"
    ws.append(COLUMNS)
    head = PatternFill("solid", fgColor="174A3A")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head

    rows = []
    for s in samples:
        step = f"{s['step']} – {s['step_name']}" if s["step_name"] else s["step"]
        rows.append((s["ts"], [
            run_id, run["protocol_name"], run["protocol_version"], _ts(s["ts"]), _hms(s["elapsed_s"]),
            step, s["actual"], s["target"], s["deviation"], s["cooling_rate"], s["humidity"],
            s["status"], "Measurement", None]))
    for e in events:
        details = e["message"]
        rows.append((e["ts"], [
            run_id, run["protocol_name"], run["protocol_version"], _ts(e["ts"]), _hms(e["elapsed_s"]),
            None, None, None, None, None, None, None, e["type"], details]))
    rows.sort(key=lambda r: r[0])           # ISO timestamps sort chronologically
    for _, r in rows:
        ws.append(r)
    for col in ("D",):
        for c in ws[col][1:]:
            c.number_format = "yyyy-mm-dd hh:mm:ss"
    for i, w in enumerate([20, 30, 10, 20, 12, 28, 14, 14, 14, 14, 14, 12, 18, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Run Summary")
    summary = [
        ("Run ID", run_id), ("Protocol", run["protocol_name"]), ("Protocol version", run["protocol_version"]),
        ("Started", run["started_at"]), ("Ended", run["ended_at"]), ("Final status", run["final_status"] or "In progress"),
        ("Paused time (s)", run["paused_total_s"]), ("Measurements", len(samples)), ("Events", len(events)),
    ]
    tol = (run["protocol_snapshot"] or {}).get("tolerance")
    summary.append(("Tolerance (°C)", tol if tol is not None else "not configured"))
    for k, v in summary:
        info.append([k, v])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 40
    for c in info["A"]:
        c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
